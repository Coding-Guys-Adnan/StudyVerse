import io
import re
import docx
import PyPDF2
from datetime import datetime, timezone
from fastapi import HTTPException, UploadFile, status
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from app.models.academic import SyllabusDocument, Syllabus
from app.models.student import Student
from app.schemas.extended import SyllabusDocumentResponse, SyllabusDocumentUploadResponse

async def _verify_student_ownership(
    db: AsyncSession, teacher_id: str, student_id: str
) -> Student:
    result = await db.execute(
        select(Student).where(
            Student.id == student_id,
            Student.teacher_id == teacher_id,
        )
    )
    student = result.scalar_one_or_none()
    if not student:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Student not found",
        )
    return student

def extract_text_from_pdf(file_bytes: bytes) -> str:
    pdf_file = io.BytesIO(file_bytes)
    reader = PyPDF2.PdfReader(pdf_file)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def extract_text_from_docx(file_bytes: bytes) -> str:
    docx_file = io.BytesIO(file_bytes)
    doc = docx.Document(docx_file)
    text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
    return text

def parse_chapters_from_text(text: str, filename: str) -> list[dict]:
    # Guess subject from filename
    subject = "General"
    name_clean = filename.lower()
    if "math" in name_clean or "algebra" in name_clean or "geometry" in name_clean:
        subject = "Mathematics"
    elif "science" in name_clean or "physics" in name_clean or "chemistry" in name_clean or "biology" in name_clean:
        subject = "Science"
    elif "english" in name_clean:
        subject = "English"
    elif "history" in name_clean or "geography" in name_clean or "civics" in name_clean or "social" in name_clean:
        subject = "Social Studies"

    lines = text.split("\n")
    chapters = []
    
    # Simple regex to look for chapters/units/topics
    # E.g. "Chapter 1: Real Numbers", "Ch 2 - Polynomials", "1. Introduction to Physics"
    pattern = re.compile(
        r'^\s*(?:chapter|ch|unit|topic|sec|section)?\s*(\d+[\.\-\d]*)\s*[:\-\.]?\s*(.+)$',
        re.IGNORECASE
    )
    
    for line in lines:
        line = line.strip()
        if not line or len(line) < 4:
            continue
        
        # Check match
        match = pattern.match(line)
        if match:
            ch_num = match.group(1).strip()
            ch_title = match.group(2).strip()
            if len(ch_title) > 2:
                chapters.append({
                    "subject": subject,
                    "chapter": f"Chapter {ch_num}: {ch_title}"
                })
        elif line.lower().startswith("chapter") or line.lower().startswith("unit"):
            # Fallback for lines like "Chapter One" or "Chapter: Real Numbers"
            chapters.append({
                "subject": subject,
                "chapter": line
            })
            
    # Fallback if no specific chapters found, return first few lines as general topics
    if not chapters:
        valid_lines = [l.strip() for l in lines if len(l.strip()) > 10][:8]
        for idx, line in enumerate(valid_lines, 1):
            chapters.append({
                "subject": subject,
                "chapter": f"Topic {idx}: {line[:60]}"
            })
            
    # Deduplicate chapters
    seen = set()
    deduped = []
    for ch in chapters:
        key = (ch["subject"], ch["chapter"])
        if key not in seen:
            seen.add(key)
            deduped.append(ch)
            
    return deduped

async def upload_syllabus_doc(
    db: AsyncSession, teacher_id: str, student_id: str, file: UploadFile
) -> SyllabusDocumentUploadResponse:
    await _verify_student_ownership(db, teacher_id, student_id)
    
    file_bytes = await file.read()
    filename = file.filename or "syllabus_document"
    
    if filename.lower().endswith(".pdf"):
        extracted_text = extract_text_from_pdf(file_bytes)
    elif filename.lower().endswith(".docx"):
        extracted_text = extract_text_from_docx(file_bytes)
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Only PDF and DOCX are supported."
        )
        
    doc = SyllabusDocument(
        student_id=student_id,
        file_name=filename,
        extracted_text=extracted_text,
    )
    db.add(doc)
    await db.flush()
    await db.refresh(doc)
    
    chapters = parse_chapters_from_text(extracted_text, filename)
    
    # Transform response
    doc_res = SyllabusDocumentResponse(
        id=doc.id,
        student_id=doc.student_id,
        file_name=doc.file_name,
        extracted_text=doc.extracted_text,
        uploaded_at=doc.uploaded_at.isoformat()
    )
    
    return SyllabusDocumentUploadResponse(
        document=doc_res,
        extracted_chapters=chapters
    )

async def list_syllabus_docs(
    db: AsyncSession, teacher_id: str, student_id: str
) -> list[SyllabusDocumentResponse]:
    await _verify_student_ownership(db, teacher_id, student_id)
    result = await db.execute(
        select(SyllabusDocument).where(
            SyllabusDocument.student_id == student_id
        ).order_by(SyllabusDocument.uploaded_at.desc())
    )
    docs = result.scalars().all()
    return [
        SyllabusDocumentResponse(
            id=d.id,
            student_id=d.student_id,
            file_name=d.file_name,
            extracted_text=d.extracted_text,
            uploaded_at=d.uploaded_at.isoformat()
        )
        for d in docs
    ]
